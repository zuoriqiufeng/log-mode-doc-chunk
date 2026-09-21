"""Excel (.xlsx / .xls) 提取器"""
from extractors.base import DocumentExtractor


class ExcelExtractor(DocumentExtractor):
    def extract(self, file_path: str) -> tuple[str, dict, list]:
        import os

        ext = os.path.splitext(file_path)[1].lower()
        metadata: dict = {}
        text_parts: list[str] = []

        if ext == ".xlsx":
            self._extract_xlsx(file_path, text_parts, metadata)
        elif ext == ".xls":
            self._extract_xls(file_path, text_parts, metadata)

        return "\n\n".join(text_parts), metadata, []

    def _extract_xlsx(self, file_path: str, text_parts: list[str], metadata: dict):
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        metadata["sheets"] = wb.sheetnames
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            text_parts.append(f"=== Sheet: {sheet_name} ===")
            for row in rows:
                cleaned = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cleaned):
                    text_parts.append(" | ".join(cleaned))
        wb.close()

    def _extract_xls(self, file_path: str, text_parts: list[str], metadata: dict):
        import xlrd
        wb = xlrd.open_workbook(file_path)
        metadata["sheets"] = wb.sheet_names()
        for sheet_name in wb.sheet_names():
            ws = wb.sheet_by_name(sheet_name)
            if ws.nrows == 0:
                continue
            text_parts.append(f"=== Sheet: {sheet_name} ===")
            for row_idx in range(ws.nrows):
                row_values = ws.row_values(row_idx)
                cleaned = [str(c) if c != "" else "" for c in row_values]
                if any(c.strip() for c in cleaned):
                    text_parts.append(" | ".join(cleaned))
