"""Word 提取器（参考 Onyx read_docx_file，使用 python-docx）"""

from __future__ import annotations

import csv
import io
import os
import re
from typing import Any

from extractors.base import DocumentExtractor


# 图片占位符，index 对应 images 列表中的顺序
_IMAGE_PLACEHOLDER = "【PIC:{index}】"

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class _PseudoParagraph:
    """合并单元格导致 row.cells 失败时，用 XML <w:p> 包装一个极简段落对象。"""

    def __init__(self, p_xml):
        self._p = p_xml

    @property
    def text(self) -> str:
        return "".join(self._p.itertext())


class DocxExtractor(DocumentExtractor):
    def extract(
        self, file_path: str
    ) -> tuple[str, dict[str, Any], list[tuple[bytes, str]]]:
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)

        # 提取文本，并在图片位置插入占位符
        images: list[tuple[bytes, str]] = []
        paragraphs = []
        for p in self._iter_paragraphs(doc):
            text = p.text.strip()
            img_indices = self._collect_image_placeholders(p, doc, images)
            if img_indices:
                placeholders = "\n\n".join(
                    _IMAGE_PLACEHOLDER.format(index=idx) for idx in img_indices
                )
                text = f"{text}\n\n{placeholders}" if text else placeholders
            if text:
                paragraphs.append(text)

        # 提取表格文本（转为 CSV 格式），单元格内图片占位符已在上面处理
        table_texts = self._extract_table_texts(doc)

        all_text = "\n\n".join(paragraphs)
        if table_texts:
            all_text += "\n\n[TABLES]\n\n" + "\n\n".join(table_texts)

        # 元数据
        metadata: dict[str, Any] = {}
        core_props = doc.core_properties
        if core_props.title:
            metadata["Title"] = core_props.title
        if core_props.author:
            metadata["Author"] = core_props.author

        return all_text, metadata, images

    def _iter_paragraphs(self, doc):
        """遍历正文段落和表格单元格中的段落。"""
        for p in doc.paragraphs:
            yield p
        for table in doc.tables:
            for row in table.rows:
                try:
                    cells = list(row.cells)
                except ValueError:
                    # 合并单元格导致 grid_offset 错误时，直接遍历该行所有 <w:p>
                    for p_xml in row._tr.findall(f"{_W_NS}tc/{_W_NS}p"):
                        yield _PseudoParagraph(p_xml)
                    continue
                for cell in cells:
                    for p in cell.paragraphs:
                        yield p

    def _iter_cells(self, row):
        """安全遍历表格行中的单元格，避免合并单元格导致的 grid_offset 错误。"""
        try:
            return list(row.cells)
        except ValueError:
            # 合并单元格时回退到直接遍历 XML 的 <w:tc> 元素
            return row._tr.findall(f"{_W_NS}tc")

    def _extract_table_texts(self, doc):
        """提取所有表格为 CSV 文本，兼容合并单元格。"""
        table_texts = []

        for table in doc.tables:
            rows = []
            for row in table.rows:
                # 直接遍历 XML 单元格，不依赖 grid_offset 计算
                tcs = row._tr.findall(f"{_W_NS}tc")
                cell_texts = []
                for tc in tcs:
                    cell_text = "".join(tc.itertext()).strip()
                    cell_texts.append(cell_text)
                if cell_texts:
                    rows.append(cell_texts)
            if rows:
                output = io.StringIO()
                writer = csv.writer(output)
                writer.writerows(rows)
                table_texts.append(output.getvalue())
        return table_texts

    def _collect_image_placeholders(
        self,
        paragraph,
        doc,
        images: list[tuple[bytes, str]],
    ) -> list[int]:
        """从段落 XML 中检测内嵌图片，加入 images 列表，返回占位符 index 列表。"""
        xml = paragraph._p.xml
        # 匹配 r:embed="rIdX"，这是 docx 中图片引用的典型方式
        embed_ids = re.findall(r'r:embed="([^"]+)"', xml)

        indices: list[int] = []
        seen: set[str] = set()
        for rId in embed_ids:
            if rId in seen:
                continue
            seen.add(rId)
            rel = doc.part.rels.get(rId)
            if not rel or "image" not in rel.reltype:
                continue
            try:
                img_data = rel.target_part.blob
                img_name = os.path.basename(rel.target_part.partname)
                idx = len(images)
                images.append((img_data, img_name))
                indices.append(idx)
            except Exception:
                continue
        return indices
